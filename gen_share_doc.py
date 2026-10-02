"""生成技术分享 Word 文档"""
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ── 全局样式 ──
style = doc.styles["Normal"]
style.font.name = "微软雅黑"
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

for level in range(1, 4):
    hs = doc.styles[f"Heading {level}"]
    hs.font.name = "微软雅黑"
    hs.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    hs.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)


def add_heading(text: str, level: int = 1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    return h


def add_code(text: str, lang: str = ""):
    """代码块：灰色底 + 等宽字体"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Cm(0.5)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)
    # 灰色背景
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F2F2F2")
    shading.set(qn("w:val"), "clear")
    p.paragraph_format.element.get_or_add_pPr().append(shading)
    return p


def add_table(headers: list[str], rows: list[list[str]]):
    """添加表格"""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for para in cell.paragraphs:
            for run in para.runs:
                run.bold = True
                run.font.size = Pt(10)
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = val
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(9.5)
    return table


def add_knowledge(text: str):
    """额外知识点（高亮引用块）"""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.font.name = "微软雅黑"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x37, 0x56, 0x23)
    run.italic = True


# ═══════════════════════════════════════════════════════════════
# 封面
# ═══════════════════════════════════════════════════════════════
for _ in range(4):
    doc.add_paragraph()

title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title_p.add_run("KamaClaude 二次开发技术分享")
run.bold = True
run.font.size = Pt(26)
run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

sub_p = doc.add_paragraph()
sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = sub_p.add_run("从开源 Agent 框架到 Windows 原生适配 + 执行持久化 + 智能摘要的完整实战")
run.font.size = Pt(13)
run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

doc.add_paragraph()
meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = meta.add_run("涵盖 S8 执行持久化 / 跨重启恢复 / Windows Shell Worker  ·  S9 统一摘要系统 / Tool Result 截断")
run.font.size = Pt(11)
run.font.color.rgb = RGBColor(0x40, 0x40, 0x40)


doc.add_page_break()

# ═══════════════════════════════════════════════════════════════
# 目录占位（手动更新）
# ═══════════════════════════════════════════════════════════════
add_heading("目录", level=1)
toc = [
    "模块一：SQLite 执行持久化层（S8）",
    "模块二：跨重启恢复机制（S8）",
    "模块三：Windows Shell Worker + Job Object（S8 Windows 适配）",
    "模块四：统一摘要系统（S9）",
    "模块五：Tool Result 智能截断（S9 Step 1）",
    "总结：Windows 二次开发避坑 Checklist",
]
for t in toc:
    doc.add_paragraph(t, style="List Number")

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════
# 模块一
# ═══════════════════════════════════════════════════════════════
add_heading("模块一：SQLite 执行持久化层（S8）", level=1)

add_heading("功能一句话", level=2)
doc.add_paragraph("把原来散落的 JSON 运行记录替换成 SQLite 十二张表（sessions / runs / messages / tool_calls / checkpoints / summaries / daemon_state…），daemon 崩了或电脑重启后能从断点续跑。")

add_heading("为什么选 SQLite", level=2)
doc.add_paragraph(
    "LangGraph 官方提供 SqliteSaver、CrewAI 内置 SqliteProvider，2026 年 Zylos Research 调研报告明确指出："
    "AI Agent 生态已经收敛到 SQLite 作为默认操作数据库。它零配置、单文件、ACID 事务、WAL 模式下 2 万+ 写/秒，"
    "完美匹配单机本地 agent 的需求——比 Postgres 轻量，比 JSON 可靠。"
)

add_heading("关键代码", level=2)

p = doc.add_paragraph()
run = p.add_run("1. Schema 版本常量")
run.bold = True
add_code("""# src/kama_claude/core/session/execution.py:25
SCHEMA_VERSION = 4  # 每次表结构变更 +1""", "python")

p = doc.add_paragraph()
run = p.add_run("2. 幂等迁移（新库从 v2 跳 v4 和旧库 v3→v4 都安全）")
run.bold = True
add_code("""# _migrate() 内统一入口
def _add_column(self, table, col, defn):
    cols = [r[1] for r in self._conn.execute(f"PRAGMA table_info({table})")]
    if col not in cols:
        self._conn.execute(f"ALTER TABLE {table} ADD COLUMN {col} {defn}")

# v3→v4：summaries 表加 summary_kind 列
self._add_column("summaries", "summary_kind", "TEXT NOT NULL DEFAULT 'handoff'")
self._conn.execute("PRAGMA user_version=4")""", "python")

p = doc.add_paragraph()
run = p.add_run("3. Checkpoint CAS 乐观锁")
run.bold = True
add_code("""def commit_summary(self, sid, run_id, *, summary_text,
                   summary_kind: str = "handoff") -> str:
    # 校验 checkpoint_version 没有被其他线程抢先写入
    row = self._conn.execute(
        "SELECT checkpoint_version FROM checkpoints WHERE session_id = ?",
        (sid,)
    ).fetchone()
    if row and row[0] != expected_version:
        raise StorageError("checkpoint conflict")
    # 校验通过才 INSERT summaries + UPDATE checkpoints""", "python")

add_heading("踩坑记录", level=2)
add_table(
    ["#", "踩的坑", "解决方案"],
    [
        ["1", "沙箱不允许操作 .git/index.lock，git add/commit 从 Trae 里跑会炸",
         "所有 git 操作在 PyCharm 终端或 Git Bash 里执行"],
        ["2", "SQLite WAL 模式多线程写冲突 database is locked",
         "单 daemon 进程 + _migrate() 用 BEGIN IMMEDIATE 事务锁住"],
        ["3", "daemon 占用 .exe 导致 editable install 失败",
         "用 PYTHONPATH=src 代替 pip install -e ."],
        ["4", "Schema 迁移只写 v3→v4 分支，新库从 v2 跳 v4 会漏列",
         "统一用 _add_column 幂等方法，所有版本走同一个入口"],
    ]
)

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════
# 模块二
# ═══════════════════════════════════════════════════════════════
add_heading("模块二：跨重启恢复机制（S8）", level=1)

add_heading("功能一句话", level=2)
doc.add_paragraph("daemon 崩溃后，kama resume 从最近 checkpoint 续跑，已经调过的工具（shell 命令、文件写入、MCP 调用）不重放，直接复用结果。")

add_heading("副作用分级策略", level=2)
add_table(
    ["工具类型", "示例", "重放策略"],
    [
        ["纯读（pure）", "echo hello、cat file、搜索", "可安全重放"],
        ["有副作用（side_effect）", "git push、写文件、发请求", "结果 unknown 时自动暂停等人审"],
    ]
)

add_heading("关键代码（概念示意）", level=2)
add_code("""# SessionManager._execute()
def _execute(self, sid):
    checkpoint = self._store.load_latest_checkpoint(sid)
    if checkpoint and checkpoint.status == "interrupted":
        # 从 checkpoint 恢复 messages + tool_results
        restored_messages = checkpoint.messages
        # 标记已完成的 tool_call，AgentLoop 跳过执行直接复用
        completed = self._store.load_completed_tool_calls(checkpoint.run_id)
    # 正常启动 agent loop，传入恢复状态""", "python")

add_heading("踩坑记录", level=2)
add_table(
    ["#", "踩的坑", "解决方案"],
    [
        ["1", "不知道哪些工具能重放、哪些不能",
         "给工具加副作用分级，side_effect 类在 unknown 时自动暂停"],
        ["2", "daemon 强杀时工具正在执行（如 ping -t），SQLite 里是 dispatching",
         "dispatching 统一归为 unknown → needs_review，TUI 弹提示"],
        ["3", "E2E 测试硬编码模型名 claude-sonnet-4-6，本地 .env 配的是 DashScope",
         "改成 os.environ.get() 读配置，不硬编码"],
    ]
)

add_knowledge("20 种崩溃场景全部通过故障矩阵验证：daemon kill -9、断电、网络断开、MCP 超时……恢复点正确，工具调用不重放，最终状态一致。")

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════
# 模块三
# ═══════════════════════════════════════════════════════════════
add_heading("模块三：Windows Shell Worker + Job Object（S8 Windows 适配）", level=1)

add_heading("功能一句话", level=2)
doc.add_paragraph("Agent 执行 shell 命令时，通过独立的 Python 子进程（_shell_worker.py）跑 cmd.exe，并用 Windows Job Object 绑定父子进程生命周期——父进程挂了子进程树被系统强制清理，不会残留僵尸进程。")

add_heading("架构图（文字版）", level=2)
add_code("""Agent 调用 bash 工具
    ↓
loop.py → invoke_tool() → BashTool.invoke()
    ↓
process.py → spawn _shell_worker.py 子进程
    ↓  worker 启动后第一行先 join_job()
Windows Job Object（KillOnJobClose）
    ↓  worker 握手 KAMA_SHELL_READY → 等父进程发命令
cmd.exe /d /s /c "用户命令"
    ↓
返回结果 → 父进程若挂 → Job Object 内核自动 kill 子进程树""", "text")

add_heading("关键代码 1：Worker 入口", level=2)
add_code("""# src/kama_claude/core/tools/_shell_worker.py:1-35
def main() -> int:
    if os.name == "nt":
        join_job(sys.argv[1])     # 启动时先加入父进程的 Job
    sys.stdout.buffer.write(b"KAMA_SHELL_READY\\n")
    sys.stdout.buffer.flush()
    payload = sys.stdin.buffer.readline()     # 等父进程发命令
    command = json.loads(payload)["command"]
    if os.name == "nt":
        shell = os.path.join(os.environ["SystemRoot"],
                             "System32", "cmd.exe")
        process = subprocess.Popen(
            f'"{shell}" /d /s /c "{command}"',
            stdin=subprocess.DEVNULL
        )
    else:
        process = subprocess.Popen(
            ["/bin/sh", "-c", command],
            stdin=subprocess.DEVNULL
        )
    return process.wait()""", "python")

add_heading("关键代码 2：Job Object 核心", level=2)
add_code("""# src/kama_claude/core/tools/windows_job.py:69-83
class WindowsJob:
    def __init__(self) -> None:
        self.name = "Local\\\\KamaClaude-" + uuid.uuid4().hex
        self.handle = CreateJobObjectW(None, self.name)
        limits = _ExtendedLimits()
        # JOB_OBJECT_LIMIT_KILL_ON_JOB_CLOSE = 0x2000
        # 告诉内核：没人持有这个 Job 的句柄时，自动 kill 里面所有进程
        limits.basic.flags = 0x2000
        SetInformationJobObject(
            self.handle,
            JobObjectExtendedLimitInformation,
            byref(limits), sizeof(limits)
        )""", "python")

add_knowledge("JOB_OBJECT_LIMIT_KILL_ON_JOB_CLOSE 是关键——它把子进程树的清理工作交给了 Windows 内核，不依赖 Python 的 atexit 或 signal 处理。daemon 被 kill -9 也能干净清理。")

add_heading("踩坑记录", level=2)
add_table(
    ["#", "踩的坑", "解决方案"],
    [
        ["1", "ctypes 在 64 位 Python 上 HANDLE 被截断成 32 位",
         "手动声明所有 Win32 函数签名（argtypes + restype）"],
        ["2", "Python 启动 worker 和 worker 加入 Job 有竞态",
         "worker 进程第一行先 join_job() 再 Popen(cmd.exe)"],
        ["3", "PowerShell Core (pwsh) 用户没装就报 not recognized",
         "worker 只认 cmd.exe，不用 pwsh；agent prompt 里明确这点"],
        ["4", "cmd.exe /c 引号嵌套地狱",
         "用 f'\"{shell}\" /d /s /c \"{command}\"' 完整命令串让 cmd 自己转义"],
    ]
)

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════
# 模块四
# ═══════════════════════════════════════════════════════════════
add_heading("模块四：统一摘要系统（S9）", level=1)

add_heading("功能一句话", level=2)
doc.add_paragraph("原来两套摘要各存各的——compactor 写的（给 agent 恢复用）和 /summarize skill 生成的人类可读摘要没有关联。现在统一到同一张 summaries 表，用 summary_kind 列区分用途。")

add_heading("两套摘要的调用路径", level=2)
add_code("""compactor 自动压缩
    ↓ loop.py 触发 或 TUI /compact
Compactor.compact_persisted()
    ↓
store.commit_summary()  → summaries 表 (kind='handoff')
    ↓
agent 恢复时读取这张表

/summarize skill
    ↓ 用户手动触发
agent 执行 → 最后一条 assistant message 就是人类可读摘要
    ↓ manager.py hook
store.write_skill_summary()  → summaries 表 (kind='human_readable')
    ↓
人类查看用这张表""", "text")

add_heading("关键代码", level=2)

p = doc.add_paragraph()
run = p.add_run("Schema 变更：summaries 表加一列")
run.bold = True
add_code("""-- SQLite schema v4 新增列
ALTER TABLE summaries
  ADD COLUMN summary_kind
  TEXT NOT NULL DEFAULT 'handoff';

-- 取值
-- 'handoff'          → compactor 写的，agent 恢复用
-- 'human_readable'   → /summarize skill 写的，人类看""", "sql")

p = doc.add_paragraph()
run = p.add_run("write_skill_summary 为什么不走 CAS 校验")
run.bold = True
add_code("""def write_skill_summary(self, sid, run_id, *, summary_text):
    # 不走 checkpoint CAS，不修改 summary_ref
    # source_checkpoint_version 写 0，纯存储用途
    self._conn.execute(
        "INSERT INTO summaries(... , summary_kind, ...) "
        "VALUES(... , 'human_readable', ...)"
    )""", "python")

add_heading("踩坑记录", level=2)
add_table(
    ["#", "踩的坑", "解决方案"],
    [
        ["1", "新库 CREATE TABLE 和旧库 ALTER TABLE 迁移路径不同",
         "_migrate() 统一入口，_add_column() 先检查列是否存在"],
        ["2", "compactor 写完要不要派生人类可读版（方案 A vs 方案 B）",
         "选方案 B 手动触发，先落地核心链路，自动化后续迭代"],
        ["3", "PRAGMA user_version 设置位置",
         "必须在 BEGIN IMMEDIATE 事务内部执行"],
    ]
)

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════
# 模块五
# ═══════════════════════════════════════════════════════════════
add_heading("模块五：Tool Result 智能截断（S9 Step 1）", level=1)

add_heading("功能一句话", level=2)
doc.add_paragraph("tool_result_limit=8000 / tool_result_keep=4000 这两个配置原来就存在但没人接线——loop.py 里两处工具调用结果追加进 context 前，超过 8000 字符自动截断为'前 4000 + 中间省略标记 + 后 2000'。")

add_heading("两层截断链路", level=2)
add_code("""原始 shell 输出（可能 200KB+）
    ↓ 第一层：bash worker 硬限制
64KB 截断 + "\\n[truncated]"     ← process.py:_MAX_OUTPUT_BYTES
    ↓ 第二层：loop.py 追加前
~6000 字符截断 + "... (truncated, full length XXXXX chars) ..."
    ↓
agent LLM 收到这个""", "text")

add_knowledge("两层截断都保留：worker 截大输出防止 Python 爆内存，loop.py 截中间量防止撑爆 context window。想精准触发第二层，用输出在 8KB~64KB 之间的命令。")

add_heading("关键代码", level=2)
add_code("""# src/kama_claude/core/loop.py
class AgentLoop:
    def __init__(self, ...,
                 tool_result_limit: int = 8000,
                 tool_result_keep: int = 4000):
        self.tool_result_limit = tool_result_limit
        self.tool_result_keep = tool_result_keep

    def _truncate_tool_result(self, text: str,
                              limit: int = None,
                              keep: int = None) -> str:
        limit = limit or self.tool_result_limit
        keep = keep or self.tool_result_keep
        if len(text) <= limit:
            return text
        tail = max(keep // 2, 1000)
        return (
            f"{text[:keep]}\\n"
            f"... (truncated, full length {len(text)} chars) ...\\n"
            f"{text[-tail:]}"
        )

    def invoke_tool(self, tool_call):
        result = self.tools[tool_call.name].invoke(...)
        # 追加进 context 前先截断
        truncated = self._truncate_tool_result(
            result.content, self.tool_result_limit, self.tool_result_keep
        )
        self.messages.append({
            "role": "tool",
            "content": truncated
        })""", "python")

add_heading("踩坑记录", level=2)
add_table(
    ["#", "踩的坑", "解决方案"],
    [
        ["1", "两层截断叠加：worker 64KB + loop.py 8KB",
         "两层都保留，各司其职"],
        ["2", "怎么验证截断有没有生效",
         "看 tool_result 里有没有 ... (truncated, full length ...) 标记"],
        ["3", "agent 自己加 pwsh -Command 前缀但 worker 只认 cmd.exe",
         "在 system prompt 里明确：shell 工具通过 cmd.exe 执行"],
    ]
)

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════
# 总结
# ═══════════════════════════════════════════════════════════════
add_heading("总结：Windows 二次开发避坑 Checklist", level=1)

add_table(
    ["维度", "踩的坑", "对策"],
    [
        ["Shell 执行", "pwsh 不存在 / 引号嵌套 / 进程残留",
         "只认 SystemRoot\\System32\\cmd.exe；Worker + Job Object 双层隔离"],
        ["Git 操作", "沙箱锁 .git/index.lock；PowerShell 不认 &&",
         "用 Git Bash 或 PyCharm 终端；bash 里 &&，PowerShell 里 ;"],
        ["路径处理", "硬编码绝对路径；Git Bash 路径格式",
         "全部用 Path + expanduser()；Git Bash 用 /f/exercise/..."],
        ["数据库", "WAL 锁 / 迁移幂等 / schema 版本",
         "BEGIN IMMEDIATE 事务；_add_column() 先检查再 ADD；每次改 schema bump 版本号"],
        ["测试", "硬编码模型名；daemon 占端口",
         "读 .env 配置；跑矩阵前先 kama-core stop"],
        ["安全", ".env 泄露；测试日志里绝对路径",
         ".gitignore 排除 .env / docs/*/evidence/；零硬编码路径"],
    ]
)

doc.add_paragraph()
add_knowledge("关键启示：Agent 框架的真正难点不在于让 LLM 输出正确的 tool_call JSON，而在于——当工具真的执行了外部命令、真的写了文件、真的发了请求之后，如何让这一切变得可恢复、可审计、可安全暂停。")

# ═══════════════════════════════════════════════════════════════
# 保存
# ═══════════════════════════════════════════════════════════════
out = r"F:\exercise\2026\kamaclaude - v1\KamaClaude\docs\s9\KamaClaude_二次开发技术分享.docx"
doc.save(out)
print(f"✅ 已生成：{out}")
